from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'LiteUp-Security-Briefing.pptx'

FOREST = RGBColor(16, 40, 32)
GREEN = RGBColor(20, 125, 89)
MINT = RGBColor(217, 243, 232)
PAPER = RGBColor(244, 247, 243)
INK = RGBColor(23, 37, 30)
MUTED = RGBColor(91, 111, 100)
WHITE = RGBColor(255, 255, 255)
AMBER = RGBColor(237, 170, 66)
RED = RGBColor(171, 59, 58)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def rect(slide, x, y, w, h, color, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False, font='Aptos', align=None, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    if align is not None:
        paragraph.alignment = align
    return box


def base_slide(kicker, title, number):
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PAPER
    rect(slide, 0, 0, 13.333, 0.13, GREEN)
    text(slide, kicker.upper(), 0.72, 0.48, 11.8, 0.28, 10, GREEN, True)
    text(slide, title, 0.72, 0.85, 11.9, 0.7, 28, FOREST, True)
    rect(slide, 0.72, 7.08, 11.9, 0.015, RGBColor(214, 224, 217))
    text(slide, 'LITEUP  |  LOCAL SECURITY BRIEFING', 0.72, 7.16, 8.0, 0.18, 8, MUTED, True)
    text(slide, f'{number:02d}', 12.05, 7.13, 0.55, 0.22, 9, GREEN, True, align=PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, heading, body, accent=GREEN, body_size=14):
    rect(slide, x, y, w, h, WHITE, True)
    rect(slide, x, y, 0.08, h, accent)
    text(slide, heading, x + 0.24, y + 0.2, w - 0.46, 0.34, 16, FOREST, True)
    text(slide, body, x + 0.24, y + 0.66, w - 0.46, h - 0.82, body_size, INK)


# 1: Cover
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = FOREST
photo = ROOT / 'asset' / 'images' / 'liteup-utility-hero.jpg'
if photo.exists():
    slide.shapes.add_picture(str(photo), Inches(6.9), Inches(0), width=Inches(6.433), height=Inches(7.5))
rect(slide, 6.65, 0, 0.55, 7.5, FOREST)
rect(slide, 0.75, 0.85, 0.68, 0.08, AMBER)
text(slide, 'LITEUP  /  SECURITY', 0.75, 1.18, 5.7, 0.35, 12, MINT, True)
text(slide, 'Safer utility\npayments, made\neasy to explain.', 0.75, 1.82, 5.85, 2.55, 34, WHITE, True)
text(slide, 'A short guide to common web risks, safe local checks, and what LiteUp does to reduce them.', 0.78, 4.78, 5.5, 1.0, 18, RGBColor(218, 230, 222))
text(slide, 'Security briefing  |  October 2026', 0.78, 6.75, 5.5, 0.25, 10, MINT, True)

# 2: Explain the model
slide = base_slide('Start here', 'A simple way to think about web security', 2)
card(slide, 0.72, 1.82, 3.82, 3.9, '1  Check the request', 'A browser is not a trusted source. Validate every form on the server, even if the page already has required fields.', GREEN, 15)
card(slide, 4.76, 1.82, 3.82, 3.9, '2  Protect the action', 'Use prepared SQL, CSRF tokens, login limits, and role checks before changing accounts or money.', AMBER, 15)
card(slide, 8.80, 1.82, 3.82, 3.9, '3  Verify the result', 'Test only your localhost app and dummy accounts. A protected test should fail safely, without changing real data.', RED, 15)
rect(slide, 0.72, 6.02, 11.9, 0.66, MINT, True)
text(slide, 'These checks are for your own local LiteUp instance only. Do not test systems you do not own.', 0.98, 6.22, 11.3, 0.22, 13, FOREST, True)

# 3: Injection & inclusion
slide = base_slide('Injection', 'SQL injection and file inclusion', 3)
card(slide, 0.72, 1.76, 5.72, 2.0, 'If vulnerable', 'A crafted login could bypass the sign-in check or expose records. A request-chosen file could expose private server files.', RED, 14)
card(slide, 6.70, 1.76, 5.92, 2.0, 'LiteUp protection', 'Prepared statements bind values separately. PHP includes use fixed application paths; request values do not choose files.', AMBER, 14)
rect(slide, 0.72, 4.02, 11.9, 2.45, FOREST, True)
text(slide, 'SAFE LOCAL CHECK  |  sign-in form only', 1.0, 4.28, 10.9, 0.28, 11, MINT, True)
text(slide, "Test identity with a wrong password:\n' OR 1=1 --", 1.0, 4.68, 4.7, 0.7, 14, WHITE, True, font='Consolas')
text(slide, 'PROTECTED CODE  |  public/login.php:18\nWHERE email = ? OR phone = ?\n$stmt->bind_param("ss", $identity, $identity);\nExpected: “Invalid login details”', 1.0, 5.40, 6.25, 0.9, 10, MINT, True, font='Consolas')
sql_evidence = ROOT / 'docs' / 'evidence-sql-injection.png'
if sql_evidence.exists():
    rect(slide, 7.85, 4.16, 4.18, 2.35, WHITE, True)
    slide.shapes.add_picture(str(sql_evidence), Inches(7.9), Inches(4.21), width=Inches(4.08), height=Inches(2.295))
text(slide, 'REAL LOCAL TEST  |  Rejected; no SQL error or login.', 7.85, 6.54, 4.18, 0.18, 8, MINT, True)

# 4: CSRF and brute force
slide = base_slide('Accounts', 'CSRF and password guessing', 4)
card(slide, 0.72, 1.78, 5.72, 2.08, 'CSRF: a forged click', 'If vulnerable, a site could make your browser submit an action you did not choose. LiteUp checks a session token on POST forms.', GREEN, 14)
card(slide, 6.70, 1.78, 5.92, 2.08, 'Brute force: repeated guesses', 'If vulnerable, an attacker could keep guessing passwords. LiteUp limits attempts; the local test blocks after five.', AMBER, 14)
rect(slide, 0.72, 4.15, 11.9, 2.72, WHITE, True)
text(slide, 'SAFE LOCAL CHECKS  |  PowerShell, from the project folder', 0.98, 4.36, 6.7, 0.28, 11, GREEN, True)
text(slide, 'php .\\test_rate_limiter.php\ncurl.exe -i http://localhost:8080/liteup_app/public/logout.php\ncurl.exe -i -X POST http://localhost:8080/liteup_app/public/logout.php', 0.98, 4.78, 6.65, 1.22, 11, INK, True, font='Consolas')
text(slide, 'Expected: 1–5 allowed, then blocked (includes/security.php:67). GET = 405; POST without CSRF = 403. Token: hash_equals() at includes/security.php:37.', 0.98, 6.12, 6.65, 0.52, 10, INK)
csrf_evidence = ROOT / 'docs' / 'evidence-csrf-logout.png'
if csrf_evidence.exists():
    slide.shapes.add_picture(str(csrf_evidence), Inches(7.9), Inches(4.3), width=Inches(4.25), height=Inches(2.39))
text(slide, 'REAL LOCAL TEST  |  GET blocked with 405.', 7.9, 6.67, 4.25, 0.16, 9, GREEN, True)

# 5: Uploads, CORS, command injection
slide = base_slide('Input boundaries', 'Safe checks for uploads and request boundaries', 5)
card(slide, 0.72, 1.78, 2.86, 4.78, 'File uploads', 'If vulnerable: a script might run or private files leak.\n\nTest: upload a harmless .txt as a profile picture.\n\nProtected: rejected; image contents, size, and type checked.\nCode: profile.php:50.', GREEN, 12)
card(slide, 3.74, 1.78, 2.86, 4.78, 'File inclusion', 'If vulnerable: private server files could be exposed.\n\nTest: scan include/require paths.\n\nProtected: fixed project paths only; no request-selected file.\nCode: includes/auth.php.', AMBER, 12)
card(slide, 6.76, 1.78, 2.86, 4.78, 'CORS', 'If too open: another site may read private responses.\n\nTest locally:\ncurl -i -H "Origin: https://attacker.invalid" http://localhost:8080/liteup_app/public/login.php\n\nProtected: no allow-origin header.', GREEN, 10)
card(slide, 9.78, 1.78, 2.84, 4.78, 'Command injection', 'If vulnerable: input could trigger an OS command.\n\nScan only; do not run a payload:\nSelect-String -Path .\\public\\*.php -Pattern "shell_exec|system|exec|passthru"\n\nProtected: no command sinks found.', RED, 10)

# 6: Show the vulnerability outcome beside LiteUp's real protected outcome.
slide = base_slide('LiteUp test evidence', 'SQL injection: what the vulnerable and secured screens look like', 6)
text(slide, 'The vulnerable screen is a clearly marked illustration. The secured screen uses evidence captured from LiteUp’s local login test.', 0.78, 1.53, 11.8, 0.42, 13, MUTED)

# Vulnerable outcome: conceptual browser mock-up, not a live bypass or modified app.
rect(slide, 0.72, 2.08, 5.78, 3.95, WHITE, True)
rect(slide, 0.72, 2.08, 5.78, 0.42, RED, True)
text(slide, 'SIMULATED VULNERABLE RESULT', 0.98, 2.19, 5.1, 0.18, 10, WHITE, True)
rect(slide, 1.04, 2.75, 5.12, 2.92, PAPER, True)
text(slide, 'LiteUp sign in', 1.34, 2.98, 4.5, 0.34, 19, FOREST, True)
text(slide, 'Email', 1.34, 3.48, 4.4, 0.2, 10, MUTED, True)
rect(slide, 1.32, 3.75, 4.55, 0.45, WHITE, True)
text(slide, "' OR 1=1 --", 1.48, 3.87, 4.2, 0.18, 12, RED, True, font='Consolas')
rect(slide, 1.32, 4.42, 4.55, 0.44, RED, True)
text(slide, 'Sign in', 1.48, 4.53, 4.2, 0.2, 11, WHITE, True)
text(slide, 'If SQL is built from raw input, login may be bypassed.', 1.34, 5.12, 4.5, 0.36, 11, INK)

# Secured outcome: a real local test screenshot already stored with the project.
rect(slide, 6.82, 2.08, 5.80, 3.95, WHITE, True)
rect(slide, 6.82, 2.08, 5.80, 0.42, GREEN, True)
text(slide, 'LITEUP SECURED RESULT · REAL LOCAL EVIDENCE', 7.08, 2.19, 5.2, 0.18, 10, WHITE, True)
secure_evidence = ROOT / 'docs' / 'evidence-sql-injection.png'
if secure_evidence.exists():
    slide.shapes.add_picture(str(secure_evidence), Inches(7.12), Inches(2.72), width=Inches(5.18), height=Inches(2.92))
else:
    rect(slide, 7.12, 2.72, 5.18, 2.92, MINT, True)
    text(slide, 'Invalid login details.\nNo session was created.', 7.40, 3.55, 4.65, 0.8, 20, FOREST, True)
text(slide, 'Prepared statements keep input separate from SQL. Test identity: public/login.php.', 7.12, 5.72, 5.12, 0.24, 10, FOREST, True)
text(slide, 'Use only your localhost app and a dummy account; the vulnerable example is a visual teaching mock-up.', 0.78, 6.35, 11.8, 0.28, 11, MUTED, True)

# 7: Next steps and boundaries
slide = base_slide('What remains', 'Good code helps; production also needs good operations', 7)
card(slide, 0.72, 1.78, 5.72, 3.85, 'In the application now', '• Prepared SQL · public/login.php:18\n• CSRF · includes/security.php:37\n• Login throttling · includes/security.php:67\n• Upload verification · public/profile.php:50\n• Idempotent utility/tamper payments + refunds\n• CSP blocks inline JavaScript\n• Private proof downloads and escaped output', GREEN, 12)
card(slide, 6.70, 1.78, 5.92, 3.85, 'Before a company launch', '• HTTPS with a valid certificate\n• Least-privilege production database user\n• Backups and restore drills\n• Provider APIs, settlement and reconciliation\n• Central logs, monitoring, incident plan\n• Independent penetration test and privacy review', AMBER, 14)
rect(slide, 0.72, 5.95, 11.9, 0.72, MINT, True)
text(slide, 'LiteUp uses manual operations today. Airtime/data/electricity are not automatically delivered until provider APIs are integrated.', 0.98, 6.18, 11.2, 0.24, 12, FOREST, True)

prs.save(OUT)
print(OUT)
